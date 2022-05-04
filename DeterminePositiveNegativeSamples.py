import argparse
from openpyxl import load_workbook, Workbook
from os import makedirs
from os.path import join, isdir


def readDsaFile(dsaFile=None):
    print('Reading dsaFile:' + str(dsaFile))

    # Read the excel
    xlWorkbook = load_workbook(filename=dsaFile)

    typingTab = xlWorkbook['Typing']
    mfi3000Tab = xlWorkbook['3000']

    genotypes={}

    for rowIndexRaw, row in enumerate(typingTab.iter_rows()):
        if rowIndexRaw > 0:
            transplantationId = str(row[0].value).strip()
            currentGenotypes = {}
            currentGenotypes['rA1'] = str(row[1].value).strip()
            currentGenotypes['rA2'] = str(row[2].value).strip()
            currentGenotypes['rB1'] = str(row[3].value).strip()
            currentGenotypes['rB2'] = str(row[4].value).strip()
            currentGenotypes['rC1'] = str(row[5].value).strip()
            currentGenotypes['rC2'] = str(row[6].value).strip()
            currentGenotypes['rDRB11'] = str(row[7].value).strip()
            currentGenotypes['rDRB12'] = str(row[8].value).strip()
            currentGenotypes['rDRB31'] = str(row[9].value).strip()
            currentGenotypes['rDRB32'] = str(row[10].value).strip()
            currentGenotypes['rDRB41'] = str(row[11].value).strip()
            currentGenotypes['rDRB42'] = str(row[12].value).strip()
            currentGenotypes['rDRB51'] = str(row[13].value).strip()
            currentGenotypes['rDRB52'] = str(row[14].value).strip()
            currentGenotypes['rDQB11'] = str(row[15].value).strip()
            currentGenotypes['rDQB12'] = str(row[16].value).strip()
            currentGenotypes['rDQA11'] = str(row[17].value).strip()
            currentGenotypes['rDQA12'] = str(row[18].value).strip()
            currentGenotypes['rDPB11'] = str(row[19].value).strip()
            currentGenotypes['rDPB12'] = str(row[20].value).strip()
            currentGenotypes['rDPA11'] = str(row[21].value).strip()
            currentGenotypes['rDPA12'] = str(row[22].value).strip()
            currentGenotypes['rDRB3451'] = str(row[23].value).strip()
            currentGenotypes['rDRB3452'] = str(row[24].value).strip()

            currentGenotypes['dA1'] = str(row[33].value).strip()
            currentGenotypes['dA2'] = str(row[34].value).strip()
            currentGenotypes['dB1'] = str(row[35].value).strip()
            currentGenotypes['dB2'] = str(row[36].value).strip()
            currentGenotypes['dC1'] = str(row[37].value).strip()
            currentGenotypes['dC2'] = str(row[38].value).strip()
            currentGenotypes['dDRB11'] = str(row[39].value).strip()
            currentGenotypes['dDRB12'] = str(row[40].value).strip()
            currentGenotypes['dDRB31'] = str(row[41].value).strip()
            currentGenotypes['dDRB32'] = str(row[42].value).strip()
            currentGenotypes['dDRB41'] = str(row[43].value).strip()
            currentGenotypes['dDRB42'] = str(row[44].value).strip()
            currentGenotypes['dDRB51'] = str(row[45].value).strip()
            currentGenotypes['dDRB52'] = str(row[46].value).strip()
            currentGenotypes['dDQB11'] = str(row[47].value).strip()
            currentGenotypes['dDQB12'] = str(row[48].value).strip()
            currentGenotypes['dDQA11'] = str(row[49].value).strip()
            currentGenotypes['dDQA12'] = str(row[50].value).strip()
            currentGenotypes['dDPB11'] = str(row[51].value).strip()
            currentGenotypes['dDPB12'] = str(row[52].value).strip()
            currentGenotypes['dDPA11'] = str(row[53].value).strip()
            currentGenotypes['dDPA12'] = str(row[54].value).strip()
            currentGenotypes['dDRB3451'] = str(row[55].value).strip()
            currentGenotypes['dDRB3452'] = str(row[56].value).strip()
            genotypes[transplantationId] = currentGenotypes

    return genotypes


def addTyping(typingLine=None, newGenotype=None, delimiter=','):
    if(newGenotype not in typingLine
        and '?' not in newGenotype
        and len(newGenotype.strip())>1):
        typingLine= typingLine+delimiter+newGenotype.strip()
    return typingLine


def createPirchePairLine(transplantationId=None, genotype=None, newline='\n'):
    patientLine = 'Recipient_' + str(transplantationId)
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rA1'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rA2'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rB1'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rB2'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rC1'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rC2'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB11'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB12'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB31'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB32'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB41'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB42'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB51'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB52'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDQB11'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDQB12'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDQA11'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDQA12'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDPB11'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDPB12'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDPA11'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDPA12'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB3451'])
    patientLine = addTyping(typingLine=patientLine, newGenotype=genotype['rDRB3452'])

    donorLine = 'Donor_' + str(transplantationId)
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dA1'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dA2'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dB1'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dB2'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dC1'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dC2'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB11'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB12'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB31'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB32'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB41'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB42'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB51'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB52'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDQB11'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDQB12'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDQA11'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDQA12'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDPB11'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDPB12'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDPA11'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDPA12'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB3451'])
    donorLine = addTyping(typingLine=donorLine, newGenotype=genotype['dDRB3452'])

    pirchPairLine=patientLine + newline + donorLine
    return pirchPairLine


def createPircheInputFiles(genotypes=None, outputDir=None, batchSize=100, newline='\n', delimiter=','):
    pircheDir = join(outputDir, 'PIRCHE_InputFiles')
    print('Creating PIRCHE Input Files:' + str(pircheDir))
    if not isdir(pircheDir):
        makedirs(pircheDir)

    batchIndex = 1
    currentBatchCount = 1
    outputFileName = join(pircheDir, 'PircheInputBatch_' + str(batchIndex) + '.csv')
    outputFile = open(outputFileName, 'w')

    for index, transplantationId in enumerate(list(genotypes.keys())):


        pirchePairLine = createPirchePairLine(transplantationId=transplantationId, genotype=genotypes[transplantationId])
        outputFile.write(pirchePairLine+newline)


        currentBatchCount += 1

        if(currentBatchCount>int(batchSize)):
            outputFile.close()
            batchIndex += 1
            currentBatchCount = 1
            outputFileName = join(pircheDir, 'PircheInputBatch_' + str(batchIndex) + '.csv')
            outputFile = open(outputFileName, 'w')
        else:
            outputFile.write(delimiter + newline)

    outputFile.close()

if __name__ == '__main__':
    parser = argparse.ArgumentParser()

    parser.add_argument("-v", "--verbose", help="verbose operation", action="store_true")
    parser.add_argument("-d", "--dsa", help="summary excel file for dsa", required=True)
    parser.add_argument("-o", "--output", help="output dir", required=True)

    args = parser.parse_args()
    verbose = args.verbose

    if not isdir(args.output):
        makedirs(args.output)

    genotypes = readDsaFile(dsaFile=args.dsa)

    createPircheInputFiles(genotypes=genotypes, outputDir=args.output)




